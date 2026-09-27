from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer
)
from tkinter import filedialog, messagebox
from datetime import datetime

def exportar_excel(titulo, columnas, datos):
    archivo = filedialog.asksaveasfilename(
        title="Guardar archivo Excel",
        defaultextension=".xlsx",
        filetypes=[
            ("Archivos Excel", "*.xlsx")
        ],
        initialfile=f"{titulo}.xlsx"
    )
    if not archivo:
        return
    try:
        libro = Workbook()
        hoja = libro.active
        hoja.title = titulo[:31]

 # Titulo
        hoja.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=len(columnas)
        )
        celda_titulo = hoja.cell(row=1, column=1)
        celda_titulo.value = f"SECURESYS - {titulo.upper()}"
        celda_titulo.font = Font(
            bold=True,
            size=16
        )
        celda_titulo.alignment = Alignment(
            horizontal="center"
        )

# Fecha de generacion
        hoja.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=len(columnas)
        )
        celda_fecha = hoja.cell(row=2, column=1)
        celda_fecha.value = (
            "Fecha de generacion: "
            + datetime.now().strftime("%d/%m/%Y %H:%M")
        )

# Encabezados
        fila_encabezados = 4

        for columna, nombre in enumerate(columnas, start=1):
            celda = hoja.cell(
                row=fila_encabezados,
                column=columna
            )

            celda.value = nombre
            celda.font = Font(bold=True)
            celda.alignment = Alignment(
                horizontal="center"
            )

# Datos
        for fila, registro in enumerate(
            datos,
            start=fila_encabezados + 1
        ):

            for columna, valor in enumerate(
                registro,
                start=1
            ):

                hoja.cell(
                    row=fila,
                    column=columna
                ).value = valor

# Ajustar ancho
        for numero_columna in range(1, len(columnas) + 1):

            maximo = 0

            letra = hoja.cell(
                row=4,
                column=numero_columna
            ).column_letter

            for numero_fila in range(
                1,
                hoja.max_row + 1
            ):

                celda = hoja.cell(
                    row=numero_fila,
                    column=numero_columna
                )

                if celda.value is not None:

                    longitud = len(
                        str(celda.value)
                    )

                    if longitud > maximo:
                        maximo = longitud

            hoja.column_dimensions[letra].width = min(
                maximo + 3,
                40
            )

        libro.save(archivo)

        messagebox.showinfo(
            "Exportacion",
            "Archivo Excel creado correctamente."
        )

    except Exception as e:

        messagebox.showerror(
            "Error",
            f"No se pudo crear el archivo Excel.\n\n{e}"
        )

def exportar_pdf(titulo, columnas, datos):
    archivo = filedialog.asksaveasfilename(
        title="Guardar archivo PDF",
        defaultextension=".pdf",
        filetypes=[
            ("Archivos PDF", "*.pdf")
        ],
        initialfile=f"{titulo}.pdf"
    )
    if not archivo:
        return
    try:
        documento = SimpleDocTemplate(
            archivo,
            pagesize=landscape(A4),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )
        estilos = getSampleStyleSheet()
        elementos = []
# Titulo
        titulo_pdf = Paragraph(
            f"<b>SECURESYS</b><br/>{titulo.upper()}",
            estilos["Title"]
        )
        elementos.append(titulo_pdf)
        elementos.append(Spacer(1, 15))
# Fecha
        fecha = datetime.now().strftime(
            "%d/%m/%Y %H:%M"
        )
        elementos.append(
            Paragraph(
                f"Fecha de generacion: {fecha}",
                estilos["Normal"]
            )
        )
        elementos.append(Spacer(1, 15))
# Tabla
        tabla_datos = [
            columnas
        ]
        for registro in datos:
            fila = []
            for valor in registro:
                fila.append(str(valor))
            tabla_datos.append(fila)
        tabla = Table(
            tabla_datos,
            repeatRows=1
        )
        tabla.setStyle(
            TableStyle([
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#2F5597")
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),
                (
                    "ALIGN",
                    (0, 0),
                    (-1, -1),
                    "CENTER"
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE"
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor("#EAF0F8")
                    ]
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    6
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    6
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6
                )
            ])
        )
        elementos.append(tabla)
        documento.build(elementos)
        messagebox.showinfo(
            "Exportacion",
            "Archivo PDF creado correctamente."
        )
    except Exception as e:
        messagebox.showerror(
            "Error",
            f"No se pudo crear el archivo PDF.\n\n{e}"
        )